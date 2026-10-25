from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from models.elements import LinkElement, NodeElement, DirectedNodeElement
from units import ft_to_m, ft_to_mm


ALT_TYPE_NAMES = {
    1: "Geometry",
    2: "Topology",
    3: "Active Topology",
    4: "Physical",
    21: "Initial Settings",
    100: "UDX",
}


def _autosize(ws):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        max_len = max(len("" if c.value is None else str(c.value)) for c in col)
        ws.column_dimensions[letter].width = min(max_len + 2, 55)


def _header(ws):
    for c in ws[1]:
        c.font = Font(bold=True)


def _keys(elements, attr):
    result = set()
    for e in elements:
        result.update(getattr(e, attr, {}).keys())
    return sorted(result)


def _join_ids(values):
    return ", ".join(str(v) for v in values) if values else ""


def export_excel(
    groups,
    stats_map,
    report,
    output_path,
    source_path,
    scenario_context,
    scenario_list,
):
    wb = Workbook()

    # SUMMARY: enxuto
    ws = wb.active
    ws.title = "Summary"
    ws.append(["WaterGEMS Scenario-Aware Extractor", ""])
    ws.append(["Source", str(source_path)])
    ws.append(["Selected Scenario ID", scenario_context.selected_scenario_id])
    ws.append(["Active Scenario ID", scenario_context.active_scenario_id])
    ws.append([
        "Selected scenario is active",
        scenario_context.selected_scenario_id == scenario_context.active_scenario_id,
    ])
    ws.append(["Validation errors", len(report.errors)])
    ws.append(["Validation warnings", len(report.warnings)])
    ws.append([])

    ws.append(["Element Type", "Exported", "Active", "Inactive", "Unresolved"])
    _header(ws)
    for name, stat in stats_map.items():
        ws.append([
            name,
            stat.exported,
            stat.active,
            stat.inactive,
            stat.unresolved_active_topology,
        ])
    _autosize(ws)

    # DIAGNOSTICS: agregado + problemas individuais
    ws = wb.create_sheet("Diagnostics")
    ws.append(["AGGREGATED DIAGNOSTICS"])
    ws.append([
        "Element Type",
        "Examined",
        "Exported",
        "Active",
        "Inactive",
        "Unresolved Active Topology",
        "Invalid Geometry",
        "Invalid Topology",
        "Missing Physical",
        "Total Issues",
    ])
    for c in ws[2]:
        c.font = Font(bold=True)

    for name, stat in stats_map.items():
        ws.append([
            name,
            stat.examined,
            stat.exported,
            stat.active,
            stat.inactive,
            stat.unresolved_active_topology,
            stat.invalid_geometry,
            stat.invalid_topology,
            stat.missing_physical,
            len(stat.issues),
        ])

    ws.append([])
    ws.append(["Scenario hierarchy used"])
    ws.append([" -> ".join(str(x) for x in scenario_context.scenario_chain)])

    ws.append([])
    ws.append(["ELEMENT ISSUES"])
    issue_header_row = ws.max_row + 1
    ws.append(["Element Type", "ID", "Label", "Issue", "Details"])
    for c in ws[issue_header_row]:
        c.font = Font(bold=True)

    for stat in stats_map.values():
        for issue in stat.issues:
            ws.append([
                issue.element_type,
                issue.element_id,
                issue.label,
                issue.issue,
                issue.details,
            ])

    ws.append([])
    ws.append(["VALIDATION ERRORS"])
    for message in report.errors:
        ws.append([message])

    ws.append([])
    ws.append(["VALIDATION WARNINGS"])
    for message in report.warnings:
        ws.append([message])

    _autosize(ws)

    # SCENARIOS
    ws = wb.create_sheet("Scenarios")
    ws.append(["Scenario ID", "Parent ID", "Depth", "Active", "Display"])
    _header(ws)
    for s in scenario_list:
        ws.append([s.id, s.parent_id, s.depth, s.is_active, s.display_name])
    _autosize(ws)

    # SCENARIO ALTERNATIVES
    ws = wb.create_sheet("Scenario Alternatives")
    ws.append([
        "Alternative Type ID",
        "Alternative Type",
        "Priority",
        "Source Scenario ID",
        "Alternative ID",
    ])
    _header(ws)
    for type_id, values in sorted(scenario_context.alternatives_by_type.items()):
        for priority, (scenario_id, alt_id) in enumerate(values, start=1):
            ws.append([
                type_id,
                ALT_TYPE_NAMES.get(type_id, ""),
                priority,
                scenario_id,
                alt_id,
            ])
    _autosize(ws)

    # ELEMENT SHEETS
    for group_name, elements in groups.items():
        if not elements:
            continue

        ws = wb.create_sheet(group_name[:31])
        physical_keys = _keys(elements, "physical")
        initial_keys = _keys(elements, "initial_settings")
        udx_keys = _keys(elements, "udx")

        sample = elements[0]

        headers = [
            "ID",
            "Label",
            "Type",
            "Is Active",
            "Active Topology Source Scenario ID",
            "Active Topology Alternative ID",
            "Geometry Source Scenario ID",
            "Geometry Alternative ID",
            "Topology Source Scenario ID",
            "Topology Alternative ID",
            "Physical Source Scenario IDs",
            "Physical Alternative IDs",
        ]

        if isinstance(sample, LinkElement):
            headers += [
                "Start Node ID", "Start Node Label", "Start Node Type",
                "End Node ID", "End Node Label", "End Node Type",
                "Length (ft)", "Length (m)",
            ]
        elif isinstance(sample, DirectedNodeElement):
            headers += [
                "X (ft)", "Y (ft)", "X (m)", "Y (m)",
                "Elevation (ft)", "Elevation (m)",
                "Downstream Link ID",
            ]
        elif isinstance(sample, NodeElement):
            headers += [
                "X (ft)", "Y (ft)", "X (m)", "Y (m)",
                "Elevation (ft)", "Elevation (m)",
            ]

        headers += [f"Physical::{k}" for k in physical_keys]
        headers += [f"Initial::{k}" for k in initial_keys]
        headers += [f"UDX::{k}" for k in udx_keys]

        diameter_keys = [
            k for k in physical_keys
            if "Diameter" in k or "PipeDiameter" in k
        ]
        for key in diameter_keys:
            headers.append(f"SI::{key} (mm)")

        ws.append(headers)
        _header(ws)

        for e in elements:
            row = [
                e.id,
                e.label,
                e.element_type,
                e.is_active,
                e.active_topology_source_scenario_id,
                e.active_topology_alternative_id,
                e.geometry_source_scenario_id,
                e.geometry_alternative_id,
                e.topology_source_scenario_id,
                e.topology_alternative_id,
                _join_ids(e.physical_source_scenario_ids),
                _join_ids(e.physical_alternative_ids),
            ]

            if isinstance(e, LinkElement):
                row += [
                    e.start_node_id, e.start_node_label, e.start_node_type,
                    e.end_node_id, e.end_node_label, e.end_node_type,
                    e.length_ft, ft_to_m(e.length_ft),
                ]
            elif isinstance(e, DirectedNodeElement):
                row += [
                    e.x_ft, e.y_ft, ft_to_m(e.x_ft), ft_to_m(e.y_ft),
                    e.elevation_ft, ft_to_m(e.elevation_ft),
                    e.downstream_link_id,
                ]
            elif isinstance(e, NodeElement):
                row += [
                    e.x_ft, e.y_ft, ft_to_m(e.x_ft), ft_to_m(e.y_ft),
                    e.elevation_ft, ft_to_m(e.elevation_ft),
                ]

            row += [e.physical.get(k) for k in physical_keys]
            row += [e.initial_settings.get(k) for k in initial_keys]
            row += [e.udx.get(k) for k in udx_keys]
            row += [ft_to_mm(e.physical.get(k)) for k in diameter_keys]

            ws.append(row)

        _autosize(ws)

    pipes = groups.get("Pipe", [])
    if pipes:
        ws = wb.create_sheet("Pipe Geometry")
        ws.append([
            "Pipe ID", "Pipe Label", "Vertex",
            "X (ft)", "Y (ft)", "X (m)", "Y (m)"
        ])
        _header(ws)
        for p in pipes:
            for idx, (x, y) in enumerate(p.geometry_ft, start=1):
                ws.append([
                    p.id, p.label, idx,
                    x, y, ft_to_m(x), ft_to_m(y)
                ])
        _autosize(ws)

    wb.save(Path(output_path))
