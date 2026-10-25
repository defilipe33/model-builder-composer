from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
TITLE_FILL = PatternFill("solid", fgColor="D9EAF7")


def _format_header(ws, row, columns):
    for column, value in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=column, value=value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _autosize(ws, max_width=42):
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        width = max(len("" if cell.value is None else str(cell.value)) for cell in column)
        ws.column_dimensions[letter].width = min(max(width + 2, 11), max_width)


def export_component_report(
    report, rows_by_component, columns, output_path, progress_callback=None
):
    """Write one mixed-element worksheet per topological component."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumo"
    summary.sheet_view.showGridLines = False

    summary.merge_cells("A1:D1")
    title = summary["A1"]
    title.value = "Relatório de componentes desconectados"
    title.font = Font(bold=True, size=14, color="17365D")
    title.fill = TITLE_FILL

    summary.append([])
    summary.append(["Componentes identificados", len(report["components"])])
    summary.append(["Elementos visitados", report["visited_count"]])
    summary.append(["Diagnóstico completo", "Não" if report["truncated"] else "Sim"])
    if report["truncated"]:
        summary.append(["Elementos não analisados", report["unscanned_count"]])

    summary.append([])
    summary_row = summary.max_row + 1
    _format_header(summary, summary_row, [
        "Componente", "Aba", "Elementos", "Composição por tipo"
    ])

    written = 0
    for index, component in enumerate(report["components"], start=1):
        sheet_name = "Principal" if index == 1 else f"Ilha {index:03d}"
        composition = ", ".join(
            f"{element_type}: {count}"
            for element_type, count in component["counts"].items()
        )
        summary.append([
            "Componente principal" if index == 1 else f"Ilha {index}",
            sheet_name,
            component["size"],
            composition,
        ])

        sheet = workbook.create_sheet(sheet_name[:31])
        sheet.sheet_view.showGridLines = False
        _format_header(sheet, 1, columns)
        for row in rows_by_component[index - 1]:
            sheet.append([row.get(column) for column in columns])
            written += 1
            if progress_callback is not None and written % 500 == 0:
                progress_callback(written)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, sheet.max_row)}"
        _autosize(sheet)

    summary.freeze_panes = f"A{summary_row + 1}"
    summary.auto_filter.ref = f"A{summary_row}:D{summary.max_row}"
    _autosize(summary, max_width=60)
    workbook.save(Path(output_path))
    if progress_callback is not None:
        progress_callback(written)
