from pathlib import Path

from database.connection import WaterGemsDatabase
from database.schema import SchemaInspector
from readers.common import CommonReader
from readers.scenario_aware import ScenarioAwareReaders
from scenario.scenario_reader import ScenarioReader
from scenario.resolver import AlternativeResolver

from repository.model_repository import ModelRepository
from repository.query_repository import QueryRepository
from repository.selective_extraction import SelectiveExtractionService
from repository.network_topology import NetworkTopologyResolver
from repository.dynamic_fields import DynamicFieldDiscovery
from repository.table_query import (
    TableQueryService,
    geometry_fields_for_type,
    modelbuilder_fields_for_type,
)
from repository.selection_sets import WaterGemsSelectionSetRepository
from repository.spatial_repository import SpatialRepository
from repository.topology_expansion import TopologyExpansionService
from repository.component_report_export import export_component_report

from selection.lazy_engine import LazySelectionEngine
from modelbuilder.column_config import ColumnConfig
from modelbuilder.exporter import export_modelbuilder
from modelbuilder.shapefile_exporter import export_shapefiles


class AppBackend:
    def __init__(self):
        self.source_path = None
        self.db = None
        self.schema = None
        self.common = None
        self.scenario_reader = None
        self.selection_set_repo = None

        self.context = None
        self.model_repo = None
        self.query_repo = None
        self.topology_resolver = None
        self.extraction = None
        self.engine = None
        self.selection = None
        self.column_config = None
        self.table_query = None
        self.spatial_repo = None
        self.selection_undo_stack = []
        self.selection_redo_stack = []
        self.selection_history_limit = 30

    def has_open_file(self):
        return self.db is not None and self.source_path is not None

    def has_active_scenario(self):
        return self.engine is not None and self.context is not None

    def has_selection(self):
        return self.selection is not None and self.selection.count() > 0

    def reset_scenario_state(self):
        """
        Mantém SQLite aberto, mas remove tudo que depende do cenário atual.
        """
        self.context = None
        self.model_repo = None
        self.query_repo = None
        self.topology_resolver = None
        self.extraction = None
        self.engine = None
        self.selection = None
        self.column_config = None
        self.table_query = None
        self.spatial_repo = None
        self.selection_undo_stack = []
        self.selection_redo_stack = []

    def close(self):
        if self.db is not None:
            try:
                self.db.close()
            except Exception:
                pass

        self.source_path = None
        self.db = None
        self.schema = None
        self.common = None
        self.scenario_reader = None
        self.selection_set_repo = None
        self.reset_scenario_state()

    def open_file(self, path):
        self.close()

        path = Path(path)
        if not path.exists():
            raise ValueError("Arquivo não encontrado.")
        if not path.is_file():
            raise ValueError("O caminho selecionado não é um arquivo.")
        if path.suffix.lower() != ".sqlite":
            raise ValueError("Selecione um arquivo .sqlite ou .wtg.sqlite.")

        self.source_path = path
        self.db = WaterGemsDatabase(path)
        self.schema = SchemaInspector(self.db)
        self.common = CommonReader(self.db, self.schema)
        self.scenario_reader = ScenarioReader(self.db, self.schema)
        self.selection_set_repo = WaterGemsSelectionSetRepository(
            self.db, self.schema
        )

        scenarios = self.scenario_reader.tree_order()
        if not scenarios:
            self.close()
            raise RuntimeError("Nenhum cenário válido foi encontrado neste banco.")

        return scenarios

    def activate_scenario(self, scenario_id):
        if self.db is None:
            raise RuntimeError("Abra um arquivo primeiro.")

        # Troca de cenário recria toda a camada scenario-aware e SelectionSet local.
        self.reset_scenario_state()

        resolver = AlternativeResolver(self.db, self.schema)
        self.context = resolver.build_context(int(scenario_id))
        readers = ScenarioAwareReaders(resolver)

        self.model_repo = ModelRepository(
            self.db, self.schema, self.common, resolver, self.context
        )
        self.query_repo = QueryRepository(self.model_repo)
        self.topology_resolver = NetworkTopologyResolver(
            self.db, self.schema, self.common, resolver, self.context
        )
        self.extraction = SelectiveExtractionService(
            self.db,
            self.schema,
            self.common,
            readers,
            self.context,
            self.topology_resolver,
        )
        self.engine = LazySelectionEngine(
            self.model_repo,
            self.query_repo,
            self.topology_resolver,
        )
        self.selection = self.engine.new_selection("ModelBuilder Selection")

        self.column_config = ColumnConfig()
        discovery = DynamicFieldDiscovery(self.schema)
        self.spatial_repo = SpatialRepository(
            resolver, self.context, self.schema
        )
        self.table_query = TableQueryService(
            self.model_repo,
            self.query_repo,
            resolver,
            self.context,
            self.common,
            self.schema,
            spatial_repo=self.spatial_repo,
            topology_resolver=self.topology_resolver,
            extraction=self.extraction,
        )
        self.topology_expansion = TopologyExpansionService(
            self.model_repo, self.topology_resolver
        )

        for element_type, _ in self.engine.available_types():
            self.column_config.initialize_type(
                element_type,
                discovery.fields_for_type(element_type),
            )

    def _selection_snapshot(self):
        if self.selection is None:
            return {}
        return {
            element_type: set(ids)
            for element_type, ids in self.selection.selected.items()
            if ids
        }

    def _restore_selection_snapshot(self, snapshot):
        self.selection.selected = {
            element_type: set(ids)
            for element_type, ids in snapshot.items()
            if ids
        }

    def push_selection_history(self):
        if self.selection is None:
            return

        snapshot = self._selection_snapshot()

        if (
            self.selection_undo_stack
            and self.selection_undo_stack[-1] == snapshot
        ):
            return

        self.selection_undo_stack.append(snapshot)

        if len(self.selection_undo_stack) > self.selection_history_limit:
            self.selection_undo_stack.pop(0)

        # Any new mutation invalidates the redo branch.
        self.selection_redo_stack.clear()

    def can_undo_selection(self):
        return bool(self.selection_undo_stack)

    def can_redo_selection(self):
        return bool(self.selection_redo_stack)

    def undo_selection(self):
        if self.selection is None or not self.selection_undo_stack:
            return False

        current = self._selection_snapshot()
        previous = self.selection_undo_stack.pop()

        self.selection_redo_stack.append(current)
        if len(self.selection_redo_stack) > self.selection_history_limit:
            self.selection_redo_stack.pop(0)

        self._restore_selection_snapshot(previous)
        return True

    def redo_selection(self):
        if self.selection is None or not self.selection_redo_stack:
            return False

        current = self._selection_snapshot()
        next_snapshot = self.selection_redo_stack.pop()

        self.selection_undo_stack.append(current)
        if len(self.selection_undo_stack) > self.selection_history_limit:
            self.selection_undo_stack.pop(0)

        self._restore_selection_snapshot(next_snapshot)
        return True

    def available_types(self):
        if self.engine is None:
            return []
        return self.engine.available_types()

    def preview(self, element_type, limit=200, offset=0):
        if self.engine is None:
            return []
        return self.model_repo.preview(
            self.engine.normalize_type(element_type),
            limit=limit,
            offset=offset,
        )

    def select_ids(self, element_type, ids):
        self.push_selection_history()
        return self.engine.select_ids(self.selection, element_type, ids)

    def select_many(self, element_type, ids):
        self.push_selection_history()
        return self.engine.select_ids(self.selection, element_type, ids)

    def deselect_ids(self, element_type, ids):
        self.push_selection_history()
        element_type = self.engine.normalize_type(element_type)
        removed = 0
        for element_id in ids:
            if self.selection.contains(element_type, int(element_id)):
                self.selection.remove(element_type, int(element_id))
                removed += 1
        return removed

    def deselect_many(self, element_type, ids):
        return self.deselect_ids(element_type, ids)

    def is_selected(self, element_type, element_id):
        element_type = self.engine.normalize_type(element_type)
        return self.selection.contains(element_type, int(element_id))

    def add_dependencies(self):
        self.push_selection_history()
        return self.engine.include_pipe_dependencies(self.selection)

    def add_associated_customer_meters(self):
        """Add Customer Meters whose resolved associated element is selected.

        CustomerNode_Demand_Data is scenario-aware through alternative type 20.
        The query follows the same precedence as AlternativeResolver: the first
        alternative in the selected scenario chain wins for each customer.
        """
        if self.selection is None:
            return {"added": 0, "already_selected": 0, "matched": 0}
        if not self.schema.table_exists("CustomerNode_Demand_Data"):
            return {"added": 0, "already_selected": 0, "matched": 0}

        target_ids = sorted({
            int(element_id)
            for element_type, element_ids in self.selection.selected.items()
            if element_type != "Customer Meter"
            for element_id in element_ids
        })
        if not target_ids:
            return {"added": 0, "already_selected": 0, "matched": 0}

        alternatives = self.context.alternatives(20)
        if not alternatives:
            return {"added": 0, "already_selected": 0, "matched": 0}

        # Stay safely below SQLite's usual 999-parameter limit; keeping the
        # target list chunked also scales to selections with hundreds of
        # thousands of pipes or nodes.
        chunk_size = 850
        resolved_customer_ids = set()
        higher_priority_alternatives = []

        for _scenario_id, alternative_id in alternatives:
            for start in range(0, len(target_ids), chunk_size):
                target_chunk = target_ids[start:start + chunk_size]
                target_marks = ", ".join("?" for _ in target_chunk)
                params = [alternative_id, *target_chunk]
                precedence_clause = ""
                if higher_priority_alternatives:
                    priority_marks = ", ".join(
                        "?" for _ in higher_priority_alternatives
                    )
                    precedence_clause = f"""
                      AND NOT EXISTS (
                          SELECT 1
                          FROM CustomerNode_Demand_Data higher
                          WHERE higher.DomainElementID=d.DomainElementID
                            AND higher.AlternativeID IN ({priority_marks})
                      )
                    """
                    params.extend(higher_priority_alternatives)

                rows = self.db.fetchall(
                    f"""
                    SELECT DISTINCT d.DomainElementID
                    FROM CustomerNode_Demand_Data d
                    JOIN CustomerNode c ON c.DomainElementID=d.DomainElementID
                    LEFT JOIN HMIDomainElement domain
                      ON domain.DomainElementID=d.DomainElementID
                    WHERE d.AlternativeID=?
                      AND d.Demand_AssociatedElement IN ({target_marks})
                      AND COALESCE(domain.IsDeleted, 0)=0
                      {precedence_clause}
                    """,
                    params,
                )
                resolved_customer_ids.update(
                    int(row["DomainElementID"]) for row in rows
                )
            higher_priority_alternatives.append(alternative_id)

        if not resolved_customer_ids:
            return {"added": 0, "already_selected": 0, "matched": 0}

        self.push_selection_history()
        selected = self.selection.selected.setdefault("Customer Meter", set())
        already_selected = len(selected.intersection(resolved_customer_ids))
        selected.update(resolved_customer_ids)
        return {
            "added": len(resolved_customer_ids) - already_selected,
            "already_selected": already_selected,
            "matched": len(resolved_customer_ids),
        }

    def summary(self):
        return self.engine.summary(self.selection)

    def clear_selection(self):
        if self.selection is not None:
            self.push_selection_history()
            self.selection.clear()

    def available_table_fields(self, element_type):
        element_type = self.engine.normalize_type(element_type)
        groups = self.column_config.groups(element_type)

        return {
            "core": ["Core.ID", "Core.Label", "Core.IsActive"],
            "geometry": geometry_fields_for_type(element_type),
            "modelbuilder": modelbuilder_fields_for_type(element_type),
            "physical": list(groups.get("physical", [])),
            "udx": list(groups.get("udx", [])),
            "initial": list(groups.get("initial", [])),
        }

    def filtered_ids(self, element_type, filter_def=None):
        element_type = self.engine.normalize_type(element_type)
        return self.table_query.filtered_ids(element_type, filter_def)

    def table_page(self, element_type, ids, offset, limit, visible_columns):
        element_type = self.engine.normalize_type(element_type)
        return self.table_query.page(
            element_type, ids, offset, limit, visible_columns
        )

    def expression_ids(self, element_type, expression):
        element_type = self.engine.normalize_type(element_type)
        return self.table_query.expression_ids(
            element_type, expression, self.column_config
        )

    def all_expression_rows(self, expression=None):
        result = []
        for element_type, _ in self.available_types():
            ids = (
                self.expression_ids(element_type, expression)
                if expression and expression.strip()
                else self.model_repo.root_ids(element_type)
            )
            result.extend((element_type, int(x)) for x in ids)
        return result

    def all_available_table_fields(self):
        merged = {
            "core": ["Core.Type", "Core.ID", "Core.Label", "Core.IsActive"],
            "geometry": [],
            "modelbuilder": [],
            "physical": [],
            "udx": [],
            "initial": [],
        }
        for element_type, _ in self.available_types():
            groups = self.available_table_fields(element_type)
            for group in (
                "geometry",
                "modelbuilder",
                "physical",
                "udx",
                "initial",
            ):
                for field in groups.get(group, []):
                    if field not in merged[group]:
                        merged[group].append(field)
        return merged

    def all_table_page(self, typed_ids, offset, limit, visible_columns):
        rows=[]
        for element_type, element_id in typed_ids[offset:offset+limit]:
            row={
                "Core.Type": element_type,
                "Core.ID": element_id,
                "Core.Label": self.common.label(element_id),
            }
            for field in visible_columns:
                if field not in row:
                    try:
                        row[field] = self.table_query.value_for_field(
                            element_type, element_id, field
                        )
                    except Exception:
                        row[field] = None
            rows.append(row)
        return rows

    def sample_values(self, element_type, field, limit=100, scan_all=False):
        if element_type != "__ALL__":
            return self.table_query.sample_values(
                element_type, field, limit, scan_all
            )

        values=[]; seen=set()
        for t, _ in self.available_types():
            available = sum(self.available_table_fields(t).values(), [])
            if field not in available:
                continue
            for value in self.table_query.sample_values(t, field, limit, scan_all):
                key=repr(value)
                if key not in seen:
                    seen.add(key); values.append(value)
                    if not scan_all and len(values) >= limit:
                        return values
        return values

    def set_unit_system(self, system):
        if self.table_query is not None:
            self.table_query.set_unit_system(system)

    def unit_system(self):
        if self.table_query is None:
            return "SI"
        return self.table_query.unit_system()

    def display_field_name(self, field):
        if self.table_query is None:
            return field
        if field.startswith("ModelBuilder."):
            return self.column_config.display_name(field.split(".", 1)[1])
        if field in self.column_config.profile_columns():
            return self.column_config.display_name(field)
        return self.table_query.display_field_name(field)

    @staticmethod
    def profile_column_from_table_field(field):
        """Map a visible table field to its corresponding profile key."""
        direct = {
            "Core.ID": "SourceDomainElementID",
            "Core.Label": "Label",
            "Core.IsActive": "IsActive",
            "Geometry.X": "X_m",
            "Geometry.Y": "Y_m",
        }
        if field in direct:
            return direct[field]
        if field.startswith("ModelBuilder."):
            return field.split(".", 1)[1]
        return field

    def profile_columns_from_table(self, table_columns_by_type):
        enabled = set()
        known = set(self.column_config.profile_columns())
        for fields in table_columns_by_type.values():
            for field in fields:
                column = self.profile_column_from_table_field(field)
                if column in known:
                    enabled.add(column)
        return enabled

    def table_columns_from_column_profile(self, element_type, enabled_columns=None):
        """Translate enabled export/profile columns into table fields."""
        if enabled_columns is None:
            enabled_columns = {
                column
                for current_type, _ in self.available_types()
                for column in self.column_config.get(current_type)
            }
        if element_type == "__ALL__":
            result = ["Core.Type"]
            for current_type, _ in self.available_types():
                for field in self.table_columns_from_column_profile(
                    current_type, enabled_columns
                ):
                    if field not in result:
                        result.append(field)
            return result

        available = self.available_table_fields(element_type)
        modelbuilder = set(available.get("modelbuilder", []))
        dynamic = set(
            available.get("physical", [])
            + available.get("udx", [])
            + available.get("initial", [])
        )
        result = []
        for column in self.column_config.available(element_type):
            if column not in enabled_columns:
                continue
            mapped = f"ModelBuilder.{column}"
            if mapped in modelbuilder:
                result.append(mapped)
            elif column in dynamic:
                result.append(column)
        return list(dict.fromkeys(result))

    def expand_selection(
        self,
        levels=1,
        direction="both",
        max_new=50000,
        max_visited=100000,
        stop_types=None,
        include_stop=True,
        stop_on_first=False,
    ):
        if self.selection is None or self.selection.count() == 0:
            raise ValueError("A seleção está vazia.")

        seed_pairs = self.selected_pairs()

        result = self.topology_expansion.expand(
            seed_pairs,
            levels=levels,
            direction=direction,
            max_new=max_new,
            max_visited=max_visited,
            stop_types=stop_types,
            include_stop=include_stop,
            stop_on_first=stop_on_first,
        )

        self.push_selection_history()

        for element_type, element_id in result["added_pairs"]:
            if (
                element_type in dict(self.available_types())
                and self.model_repo.has_id(element_type, element_id)
            ):
                self.selection.add(element_type, element_id)

        return result

    def select_connected_component(
        self,
        stop_types=None,
        include_stop=True,
        allowed_types=None,
        inactive_as_boundary=False,
        max_new=50000,
        max_visited=100000,
    ):
        if self.selection is None or self.selection.count() == 0:
            raise ValueError("A seleção está vazia.")

        active_cache = {}

        def is_active(element_type, element_id):
            key = (element_type, int(element_id))
            if key not in active_cache:
                try:
                    active_cache[key] = self.table_query.raw_value_for_field(
                        element_type, int(element_id), "Core.IsActive"
                    )
                except Exception:
                    # Unknown activity must not make the traversal lose part
                    # of the network; only an explicit False is a boundary.
                    active_cache[key] = None
            return active_cache[key]

        result = self.topology_expansion.connected_component(
            self.selected_pairs(),
            max_new=max_new,
            max_visited=max_visited,
            stop_types=stop_types,
            include_stop=include_stop,
            allowed_types=allowed_types,
            inactive_as_boundary=inactive_as_boundary,
            is_active=is_active if inactive_as_boundary else None,
        )

        self.push_selection_history()
        available = dict(self.available_types())
        for element_type, element_id in result["added_pairs"]:
            if (
                element_type in available
                and self.model_repo.has_id(element_type, element_id)
            ):
                self.selection.add(element_type, element_id)
        return result

    def diagnose_connected_components(
        self,
        max_visited=None,
        max_components=None,
        progress_callback=None,
    ):
        if not self.has_active_scenario():
            raise RuntimeError("Carregue um cenário antes de diagnosticar a rede.")
        types = [element_type for element_type, _ in self.available_types()]
        return self.topology_expansion.diagnose_components(
            types,
            max_visited=max_visited,
            max_components=max_components,
            progress_callback=progress_callback,
        )

    def connected_component_diagnosis_estimate(self):
        if not self.has_active_scenario():
            raise RuntimeError("Carregue um cenário antes de diagnosticar a rede.")

        counts = {
            element_type: self.model_repo.count_type(element_type)
            for element_type, _ in self.available_types()
        }
        total = sum(counts.values())
        # Conservative planning estimate. Database speed and topology density
        # vary considerably, so this is deliberately presented as approximate.
        seconds = max(10, int((total + 599) // 600))
        return {
            "total_elements": total,
            "counts": counts,
            "estimated_seconds": seconds,
        }

    def network_quality_diagnostics(self, progress_callback=None):
        if not self.has_active_scenario():
            raise RuntimeError("Carregue um cenário antes de executar diagnósticos.")
        types = [element_type for element_type, _ in self.available_types()]
        nodes = [(t, int(i)) for t in types if t != "Pipe" for i in self.model_repo.root_ids(t)]
        pipes = [("Pipe", int(i)) for i in self.model_repo.root_ids("Pipe")] if "Pipe" in types else []
        total = len(pipes) + len(nodes) * 2
        done = 0
        disconnected = []
        missing_endpoints = []
        zero_length = []
        no_geometry = []
        inactive = []
        connected_node_ids = set()
        coordinate_groups = {}
        tolerance = 0.003280839895

        for _, pipe_id in pipes:
            start_id, start_type, end_id, end_type = self.topology_expansion.pipe_endpoints(pipe_id)
            if start_id is None or end_id is None:
                missing_endpoints.append(("Pipe", pipe_id))
            if start_id is not None:
                connected_node_ids.add((start_type, int(start_id)))
            if end_id is not None:
                connected_node_ids.add((end_type, int(end_id)))
            try:
                if self.table_query.raw_value_for_field("Pipe", pipe_id, "Geometry.IsZeroLength"):
                    zero_length.append(("Pipe", pipe_id))
                if not self.table_query.raw_value_for_field("Pipe", pipe_id, "Geometry.HasGeometry"):
                    no_geometry.append(("Pipe", pipe_id))
            except Exception:
                pass
            done += 1
            if progress_callback and done % 250 == 0:
                progress_callback(done, total, "Verificando Pipes")

        for element_type, element_id in nodes:
            if (element_type, element_id) not in connected_node_ids:
                disconnected.append((element_type, element_id))
            try:
                if not self.table_query.raw_value_for_field(element_type, element_id, "Geometry.HasGeometry"):
                    no_geometry.append((element_type, element_id))
                    continue
                x = self.table_query.raw_value_for_field(element_type, element_id, "Geometry.X")
                y = self.table_query.raw_value_for_field(element_type, element_id, "Geometry.Y")
                if x is not None and y is not None:
                    key = (round(float(x) / tolerance), round(float(y) / tolerance))
                    coordinate_groups.setdefault(key, []).append((element_type, element_id))
            except Exception:
                pass
            done += 1
            if progress_callback and done % 250 == 0:
                progress_callback(done, total, "Verificando nós")

        for element_type, element_id in pipes + nodes:
            try:
                if self.table_query.raw_value_for_field(element_type, element_id, "Core.IsActive") is False:
                    inactive.append((element_type, element_id))
            except Exception:
                pass
            done += 1
            if progress_callback and done % 250 == 0:
                progress_callback(done, total, "Verificando atividade")

        coincident = [pair for group in coordinate_groups.values() if len(group) > 1 for pair in group]
        if progress_callback:
            progress_callback(total, total, "Concluindo diagnóstico")
        return {
            "Pipes sem nó inicial ou final": missing_endpoints,
            "Pipes com comprimento aproximadamente zero": zero_length,
            "Nós isolados": disconnected,
            "Nós coincidentes": coincident,
            "Elementos sem geometria": no_geometry,
            "Elementos inativos": inactive,
        }

    def export_network_quality_diagnostics(self, diagnostics, output_path, progress_callback=None):
        from openpyxl import Workbook
        from openpyxl.styles import Font

        workbook = Workbook()
        summary = workbook.active
        summary.title = "Resumo"
        summary.append(["Diagnóstico de qualidade da rede"])
        summary["A1"].font = Font(bold=True, size=14)
        summary.append([])
        summary.append(["Verificação", "Itens", "Aba"])
        for cell in summary[3]:
            cell.font = Font(bold=True)

        total = sum(len(pairs) for pairs in diagnostics.values())
        done = 0
        for index, (name, pairs) in enumerate(diagnostics.items(), start=1):
            sheet_name = f"D{index:02d}"
            summary.append([name, len(pairs), sheet_name])
            sheet = workbook.create_sheet(sheet_name)
            sheet.append(["Tipo", "ID", "Label"])
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            labels = self.common.labels([element_id for _, element_id in pairs])
            for element_type, element_id in pairs:
                sheet.append([element_type, int(element_id), labels.get(int(element_id), "")])
                done += 1
                if progress_callback and done % 500 == 0:
                    progress_callback(done, total, "Gravando Excel")
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = f"A1:C{max(1, sheet.max_row)}"
            sheet.column_dimensions["A"].width = 32
            sheet.column_dimensions["B"].width = 14
            sheet.column_dimensions["C"].width = 32
        summary.column_dimensions["A"].width = 44
        summary.column_dimensions["B"].width = 14
        summary.column_dimensions["C"].width = 12
        workbook.save(output_path)
        if progress_callback:
            progress_callback(total, total, "Excel concluído")

    def export_connected_component_report(
        self,
        report,
        output_path,
        include_all_fields=False,
        progress_callback=None,
    ):
        if not report or "components" not in report:
            raise ValueError("Não há diagnóstico de componentes para exportar.")

        component_pairs = [component.get("pairs", []) for component in report["components"]]
        total = sum(len(pairs) for pairs in component_pairs)

        if include_all_fields:
            field_groups = self.all_available_table_fields()
            columns = []
            for group_name in (
                "core", "geometry", "modelbuilder", "physical", "udx", "initial"
            ):
                for field in field_groups.get(group_name, []):
                    if field not in columns:
                        columns.append(field)
            columns = ["Core.Type"] + [
                field for field in columns if field != "Core.Type"
            ]
        else:
            # Fast default: enough to audit membership and jump back to the
            # original model without per-field scenario resolution.
            columns = ["Core.Type", "Core.ID", "Core.Label"]

        all_ids = [element_id for pairs in component_pairs for _, element_id in pairs]
        labels = self.common.labels(all_ids)
        rows_by_key = {}
        processed = 0

        if include_all_fields:
            ids_by_type = {}
            for pairs in component_pairs:
                for element_type, element_id in pairs:
                    ids_by_type.setdefault(element_type, []).append(int(element_id))

            for element_type, ids in ids_by_type.items():
                for offset in range(0, len(ids), 500):
                    batch = ids[offset:offset + 500]
                    rows = self.table_query.page(
                        element_type, batch, 0, len(batch),
                        [field for field in columns if field != "Core.Type"],
                    )
                    for row in rows:
                        row["Core.Type"] = element_type
                        rows_by_key[(element_type, int(row["Core.ID"]))] = row
                    processed += len(batch)
                    if progress_callback is not None:
                        progress_callback(processed, total, "Carregando campos")
        else:
            for pairs in component_pairs:
                for element_type, element_id in pairs:
                    rows_by_key[(element_type, int(element_id))] = {
                        "Core.Type": element_type,
                        "Core.ID": int(element_id),
                        "Core.Label": labels.get(int(element_id), ""),
                    }
                    processed += 1
                    if progress_callback is not None and processed % 500 == 0:
                        progress_callback(processed, total, "Preparando linhas")

        rows_by_component = [
            [rows_by_key[(element_type, int(element_id))] for element_type, element_id in pairs]
            for pairs in component_pairs
        ]

        def writing_progress(written):
            if progress_callback is not None:
                progress_callback(total + written, total * 2, "Gravando Excel")

        export_component_report(
            report,
            rows_by_component,
            columns,
            output_path,
            progress_callback=writing_progress,
        )

    def _pipe_fallback_geometry(self, pipe_id):
        """
        Fallback only for visualization.

        If a Pipe has no resolvable LineString WKB in the active Geometry
        alternative chain, draw a straight segment between its scenario-aware
        StartNode and StopNode coordinates.
        """
        try:
            resolved = self.topology_resolver.resolve_pipe_endpoints(
                int(pipe_id)
            )
        except Exception:
            return None

        start_id = resolved.get("start_node_id")
        stop_id = resolved.get("end_node_id")
        start_type = resolved.get("start_node_type")
        stop_type = resolved.get("end_node_type")

        if (
            start_id is None
            or stop_id is None
            or not start_type
            or not stop_type
        ):
            return None

        try:
            p1 = self.spatial_repo.geometry(start_type, int(start_id))
            p2 = self.spatial_repo.geometry(stop_type, int(stop_id))
        except Exception:
            return None

        if (
            isinstance(p1, tuple)
            and len(p1) == 2
            and isinstance(p2, tuple)
            and len(p2) == 2
        ):
            return [p1, p2]

        return None

    def map_geometry_diagnostics(self, typed_ids):
        diagnostics = []

        for element_type, element_id in typed_ids:
            geometry, source = self.spatial_repo.geometry_with_source(
                element_type,
                int(element_id),
            )

            diagnostics.append(
                {
                    "type": element_type,
                    "id": int(element_id),
                    "geometry_found": geometry is not None,
                    "table": source.get("table"),
                    "scenario_id": source.get("scenario_id"),
                    "alternative_id": source.get("alternative_id"),
                    "attempts": source.get("attempts", []),
                }
            )

        return diagnostics

    def map_geometries_report(self, typed_ids):
        requested = [
            (str(element_type), int(element_id))
            for element_type, element_id in typed_ids
        ]

        rows = self.map_geometries(requested)
        loaded = {
            (row["type"], int(row["id"]))
            for row in rows
        }

        fallback = []

        for element_type, element_id in requested:
            pair = (element_type, element_id)

            if pair in loaded or element_type != "Pipe":
                continue

            geometry = self._pipe_fallback_geometry(element_id)
            if geometry is None:
                continue

            try:
                is_active = self.table_query.raw_value_for_field(
                    element_type,
                    element_id,
                    "Core.IsActive",
                )
            except Exception:
                is_active = None

            rows.append(
                {
                    "type": element_type,
                    "id": element_id,
                    "geometry": geometry,
                    "is_active": is_active,
                    "geometry_source": "endpoint_fallback",
                }
            )
            loaded.add(pair)
            fallback.append(pair)

        missing = [
            pair for pair in requested
            if pair not in loaded
        ]

        return {
            "rows": rows,
            "requested": len(requested),
            "loaded": len(rows),
            "missing": missing,
            "fallback": fallback,
        }

    def map_geometries(self, typed_ids):
        if self.spatial_repo is None:
            return []

        rows = self.spatial_repo.geometries(typed_ids)

        for row in rows:
            try:
                row["is_active"] = self.table_query.raw_value_for_field(
                    row["type"],
                    row["id"],
                    "Core.IsActive",
                )
            except Exception:
                row["is_active"] = None

        return rows

    def apply_pair_selection(self, pairs, action):
        self.push_selection_history()

        grouped = {}
        for element_type, element_id in pairs:
            normalized = self.engine.normalize_type(element_type)
            grouped.setdefault(normalized, []).append(int(element_id))

        if action == "replace":
            self.selection.clear()

        for element_type, ids in grouped.items():
            if action in {"replace", "add"}:
                for element_id in ids:
                    if (
                        self.model_repo.has_id(element_type, element_id)
                        and self.model_repo.is_live_id(element_id)
                    ):
                        self.selection.add(element_type, element_id)
            elif action == "remove":
                for element_id in ids:
                    self.selection.remove(element_type, element_id)
            else:
                raise ValueError(f"Ação espacial inválida: {action}")

    def map_geometries_in_extent(
        self,
        typed_ids,
        extent,
        max_checks=40000,
        max_results=25000,
        max_seconds=10.0,
    ):
        if self.spatial_repo is None:
            return {
                "rows": [],
                "checked": 0,
                "truncated": False,
                "reason": None,
                "elapsed": 0.0,
            }

        result = self.spatial_repo.geometries_in_extent(
            typed_ids,
            extent,
            max_checks=max_checks,
            max_results=max_results,
            max_seconds=max_seconds,
        )

        for row in result["rows"]:
            try:
                row["is_active"] = self.table_query.raw_value_for_field(
                    row["type"],
                    row["id"],
                    "Core.IsActive",
                )
            except Exception:
                row["is_active"] = None

        return result

    def selected_pairs(self):
        if self.selection is None:
            return []
        pairs = []
        for element_type, ids in self.selection.selected.items():
            for element_id in ids:
                pairs.append((element_type, int(element_id)))
        return pairs

    # -------- WaterGEMS native Selection Sets --------
    def watergems_selection_sets(self):
        if self.selection_set_repo is None:
            return []
        return self.selection_set_repo.list_sets()

    def load_watergems_selection_set(self, selection_set_id, mode="add"):
        if not self.has_active_scenario():
            raise RuntimeError("Carregue um cenário antes de usar Selection Sets.")

        grouped, unsupported = self.selection_set_repo.group_live_members(
            selection_set_id,
            self.model_repo,
        )

        self.push_selection_history()

        if mode == "replace":
            self.selection.clear()
        elif mode != "add":
            raise ValueError("Modo deve ser add ou replace.")

        added = 0
        already = 0

        for element_type, ids in grouped.items():
            for element_id in ids:
                if self.selection.contains(element_type, element_id):
                    already += 1
                else:
                    self.selection.add(element_type, element_id)
                    added += 1

        return {
            "added": added,
            "already_selected": already,
            "unsupported": unsupported,
            "grouped_counts": {
                k: len(v) for k, v in grouped.items()
            },
        }

    def export(self, output_path, progress_callback=None):
        if self.selection is None or self.selection.count() == 0:
            raise ValueError("A seleção está vazia.")

        output_path = Path(output_path)
        if output_path.suffix.lower() != ".xlsx":
            output_path = output_path.with_suffix(".xlsx")

        if not output_path.parent.exists():
            raise ValueError("O diretório de saída não existe.")

        selected_groups = self.extraction.extract_selected(self.selection.selected)
        total = sum(len(elements) for elements in selected_groups.values())
        if progress_callback:
            progress_callback(0, total, "Extração concluída; preparando Excel")
        export_modelbuilder(
            selected_groups,
            output_path,
            selection_name=self.selection.name,
            scenario_id=self.context.selected_scenario_id,
            column_config=self.column_config,
            progress_callback=progress_callback,
        )
        return output_path

    def export_gis_layers(self, output_dir, progress_callback=None):
        if self.selection is None or self.selection.count() == 0:
            raise ValueError("A seleção está vazia.")

        output_dir = Path(output_dir)
        if output_dir.exists():
            raise ValueError("A pasta de saída já existe. Escolha outra pasta-base.")
        if not output_dir.parent.exists():
            raise ValueError("O diretório de saída não existe.")

        selected_groups = self.extraction.extract_selected(self.selection.selected)
        total = sum(len(elements) for elements in selected_groups.values())
        if progress_callback:
            progress_callback(0, total, "Extração concluída; preparando camadas GIS")
        export_shapefiles(
            selected_groups,
            output_dir,
            column_config=self.column_config,
            selection_name=self.selection.name,
            scenario_id=self.context.selected_scenario_id,
            progress_callback=progress_callback,
        )
        return output_dir
