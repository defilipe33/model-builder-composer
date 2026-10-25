from pathlib import Path
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from modelbuilder.profiles.watergems_2019 import PROFILE
from modelbuilder.column_config import ColumnConfig


def autosize(ws):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        max_len = max(len("" if c.value is None else str(c.value)) for c in col)
        ws.column_dimensions[letter].width = min(max_len + 2, 45)


def write_table(ws, rows, columns=None, display_names=None):
    if not rows:
        return

    if columns:
        headers = [c for c in columns if any(c in row for row in rows)]
    else:
        headers = []
        seen = set()
        for row in rows:
            for key in row:
                if key not in seen:
                    seen.add(key)
                    headers.append(key)

    display_names = display_names or {}
    ws.append([display_names.get(header, header) for header in headers])
    for c in ws[1]:
        c.font = Font(bold=True)

    for row in rows:
        ws.append([row.get(h) for h in headers])

    autosize(ws)


def export_modelbuilder(
    selected_groups,
    output_path,
    selection_name="Selection",
    scenario_id=None,
    include_pipe_vertices=True,
    column_config=None,
    progress_callback=None,
):
    column_config = column_config or ColumnConfig()

    wb = Workbook()
    summary = wb.active
    summary.title = "Selection Summary"

    summary.append(["ModelBuilder Export", PROFILE["name"]])
    summary.append(["Selection", selection_name])
    summary.append(["Scenario ID", scenario_id])
    summary.append([])
    summary.append(["Element Type", "Count"])
    for group_name, elements in selected_groups.items():
        summary.append([group_name, len(elements)])
    autosize(summary)

    sheet_rows = defaultdict(list)
    total = sum(len(elements) for elements in selected_groups.values())
    done = 0

    for group_name, elements in selected_groups.items():
        mapping = PROFILE["sheets"].get(group_name)
        if mapping is None:
            continue

        sheet_name, mapper = mapping
        for element in elements:
            row = mapper(element)
            row["SourceElementType"] = group_name
            sheet_rows[sheet_name].append(row)
            done += 1
            if progress_callback and done % 250 == 0:
                progress_callback(done, total, "Preparando tabelas ModelBuilder")

    for sheet_name, rows in sheet_rows.items():
        ws = wb.create_sheet(sheet_name[:31])

        # Determine which source element types feed this sheet.
        source_types = []
        for row in rows:
            t = row.get("SourceElementType")
            if t and t not in source_types:
                source_types.append(t)

        configured = []
        for t in source_types:
            for col in column_config.get(t):
                if col not in configured:
                    configured.append(col)

        display_names = {
            column: column_config.display_name(column)
            for column in configured
        }
        write_table(
            ws,
            rows,
            columns=configured or None,
            display_names=display_names,
        )

    # Reference-only geometry sheet. This is intentionally separated from the
    # clean ModelBuilder mapping tables.
    if include_pipe_vertices and selected_groups.get("Pipe"):
        ws = wb.create_sheet("PipeVertices")
        rows = []
        for pipe in selected_groups["Pipe"]:
            for vertex_no, (x_ft, y_ft) in enumerate(pipe.geometry_ft, start=1):
                rows.append({
                    "SourceElementID": f"WG_SRC_{pipe.id}",
                    "PipeLabel": pipe.label,
                    "Vertex": vertex_no,
                    "X_ft": x_ft,
                    "Y_ft": y_ft,
                    "X_m": x_ft * 0.3048,
                    "Y_m": y_ft * 0.3048,
                })
        write_table(ws, rows)

    notes = wb.create_sheet("Import Notes")
    notes.append(["Item", "Guidance"])
    notes.append([
        "SourceElementID",
        "Stable source key generated from the WaterGEMS DomainElementID. "
        "Map it to GIS-ID or keep it as a UDX for traceability.",
    ])
    notes.append([
        "Pipe connectivity",
        "Use StartNodeLabel/StopNodeLabel (or Source IDs if preferred) "
        "to associate each pipe with its endpoint nodes.",
    ])
    notes.append([
        "PipeVertices",
        "Reference-only sheet containing intermediate pipe vertices. "
        "Whether these bends can be imported directly depends on the "
        "ModelBuilder workflow/version.",
    ])
    notes.append([
        "IsActive",
        "Scenario-aware Active Topology value. It is exported as data, not used as a selection filter.",
    ])
    autosize(notes)

    wb.save(Path(output_path))
    if progress_callback:
        progress_callback(total, total, "Excel ModelBuilder concluído")
